from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

SEED = [
 ('OUTRAGE PRE WORKOUT',115,5.5,30),('THE SHIELD',100,3,100),('KREATINE MONOHYDRATE',120,3.8,80),
 ('ATOMIC-1',100,2.4,70),('AAKG PLUS',100,5,60),('AAKG',100,2.8,70),('CITRULLINE',100,3.9,50),
 ('GLUTA MINE',100,2.5,80),('BURN SHAPE',110,4,40),('L-KARNITINE',115,5,31),('IRON CUT',100,4,30),
 ('XTEND',210,5,90),('SUPER PUMP AGRESION',170,8,40),('AMINO MAX 8000',130,None,60)]
ROWS, FIRST = 60, 3          # რიგი 1: ჯამები, რიგი 2: სათაური
LAST = FIRST + ROWS - 1
D0 = 9                       # პირველი დღის სვეტი (I)
D1 = D0 + 30                 # 31-ე დღე

wb = Workbook(); ws = wb.active; ws.title = 'Supplements'
dark = PatternFill('solid', fgColor='111111'); white = Font(bold=True, color='FFFFFF')
thin = Border(bottom=Side(style='thin', color='E5E7EB'))
heads = ['ფოტო','დასახელება','თვითღირებულება ₾','ფასი მიღებაზე ₾','სულ მიღება','გაყიდული','დარჩენილი','შემოსავალი ₾']
for i,h in enumerate(heads,1):
    c = ws.cell(2,i,h); c.fill=dark; c.font=white; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
for d in range(1,32):
    c = ws.cell(2,D0+d-1,d); c.fill=PatternFill('solid',fgColor='7F1D1D'); c.font=white; c.alignment=Alignment(horizontal='center',vertical='center')
ws['A1']='🏋️ დანამატები'; ws['A1'].font=Font(bold=True,size=14); ws.merge_cells('A1:D1')
ws['E1']=f'="შემოსავალი: "&ROUND(SUM(H{FIRST}:H{LAST}),2)&" ₾"'; ws.merge_cells('E1:G1')
ws['H1']='← ჩაწერე დღის სვეტში, რამდენი გაიყიდა'
ws['E1'].font=Font(bold=True,size=12); ws['E1'].alignment=Alignment(horizontal='center')
ws.row_dimensions[1].height=26
ws.row_dimensions[2].height=38

dcols = f'{L(D0)}{{r}}:{L(D1)}{{r}}'
for i in range(ROWS):
    r = FIRST+i
    if i < len(SEED):
        n,c,p,t = SEED[i]; ws.cell(r,2,n); ws.cell(r,3,c); ws.cell(r,4,p); ws.cell(r,5,t)
    ws.cell(r,6,f'=IF($B{r}="","",SUM({dcols.format(r=r)}))')
    ws.cell(r,7,f'=IF($B{r}="","",N($E{r})-N($F{r}))')
    ws.cell(r,8,f'=IF($B{r}="","",N($F{r})*N($D{r}))')
    ws.row_dimensions[r].height=60
    for col in range(1,D1+1):
        c = ws.cell(r,col); c.border=thin; c.alignment=Alignment(horizontal='center' if col!=2 else 'left',vertical='center',wrap_text=(col==2))
    ws.cell(r,2).font=Font(bold=True)
    ws.cell(r,6).font=Font(bold=True,size=16); ws.cell(r,6).fill=PatternFill('solid',fgColor='FEF2F2')
    ws.cell(r,7).font=Font(size=13)
    for col in range(D0,D1+1): ws.cell(r,col).fill=PatternFill('solid',fgColor='F9FAFB')

widths={1:14,2:24,3:13,4:13,5:10,6:11,7:11,8:13}
for k,v in widths.items(): ws.column_dimensions[L(k)].width=v
for col in range(D0,D1+1): ws.column_dimensions[L(col)].width=5
rng=f'G{FIRST}:G{LAST}'
ws.conditional_formatting.add(rng,FormulaRule(formula=[f'AND(ISNUMBER($G{FIRST}),$G{FIRST}<=0)'],fill=PatternFill('solid',bgColor='FECACA',fgColor='FECACA'),font=Font(color='B91C1C',bold=True),stopIfTrue=True))
ws.conditional_formatting.add(rng,FormulaRule(formula=[f'AND(ISNUMBER($G{FIRST}),$G{FIRST}<=5)'],fill=PatternFill('solid',bgColor='FED7AA',fgColor='FED7AA')))
ws.freeze_panes=f'C{FIRST}'
wb.save('supplements-tracker.xlsx')
